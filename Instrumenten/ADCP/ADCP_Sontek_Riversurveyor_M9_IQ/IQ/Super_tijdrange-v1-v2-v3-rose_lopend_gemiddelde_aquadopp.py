#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Aquadopp (AQP) / Aquadopp-profiel verwerking:
- .hdr bevat metadata (starttijd, interval, cell distances, transformation matrix)
- .v1 .v2 .v3 bevatten beam-snelheden per cel (whitespace separated, geen tijdkolom)
- .csv bevat o.a. Speed#i / Dir#i per cel (controle/verwerkte data)

GUI:
- kies .hdr (basisnaam), vink cellen aan
- kies begin/eindmoment (tijdvenster) voor visualisatie + export
- herbereken en teken:
  * beam1/beam2/beam3: time series per geselecteerde cel + gemiddelde (alleen binnen tijdvenster)
  * resultante speed+dir uit .csv over geselecteerde cellen (alleen binnen tijdvenster)
  * X/Y/Z projecties via transformation matrix: per geselecteerde cel + gemiddelde (alleen binnen tijdvenster)
- export: tab-gescheiden, decimaal ',' en duizendtallen '.'
"""

import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import tempfile
import shutil
import warnings

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:  # pragma: no cover
    Document = None
    Inches = None


# ----------------------------
# Parsing helpers
# ----------------------------
def _parse_hdr_text(hdr_path: str) -> str:
    with open(hdr_path, "r", errors="ignore") as f:
        return f.read()


def _find_one(text: str, pattern: str, cast=None, default=None):
    m = re.search(pattern, text)
    if not m:
        return default
    val = m.group(1).strip()
    return cast(val) if cast else val


def _extract_sen_pressure_column(text: str) -> int:
    block = re.search(r"\[.*?\.sen\](.*?)(?:\r?\n\s*\r?\n|$)", text, flags=re.S | re.I)
    if not block:
        return 14
    block_text = block.group(1)
    line = re.search(r"^\s*(\d+)\s+Pressure\b", block_text, flags=re.M)
    if not line:
        return 14
    try:
        return int(line.group(1))
    except ValueError:
        return 14


def parse_hdr(hdr_path: str) -> dict:
    """
    Leest essentiÃ«le metadata uit .hdr:
    - n_meas, n_cells, profile_interval_sec, t0 (datetime)
    - cell_distances_m (list)
    - transformation_matrix (3x3)
    """
    text = _parse_hdr_text(hdr_path)

    n_meas = _find_one(text, r"Number of measurements\s+(\d+)", int)
    n_cells = _find_one(text, r"Number of cells\s+(\d+)", int)
    dt_sec = _find_one(text, r"Profile interval\s+([0-9.]+)\s*sec", float)

    t0_str = _find_one(
        text,
        r"(?im)^\s*(?:Time of first measurement|First measurement time|Time of first profile|Start time(?: of first measurement)?)\s*[:=]?\s*(.+?)\s*$",
        str
    )
    if not t0_str:
        raise ValueError("Kon 'Time of first measurement' niet vinden in .hdr.")
    t0 = pd.NaT
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M:%S"):
        t0 = pd.to_datetime(t0_str, format=fmt, dayfirst=True, errors="coerce")
        if not pd.isna(t0):
            break
    if pd.isna(t0):
        t0 = pd.to_datetime(t0_str, dayfirst=True, errors="coerce")
    if pd.isna(t0):
        raise ValueError(f"Kon starttijd niet parsen uit .hdr: '{t0_str}'.")

    # cell center distances block
    cell_distances = []
    block = re.search(
        r"Current profile cell center distance from head \(m\)\s*\n[-]+\s*\n(.*?)(?:\n\s*\n|Data file format)",
        text,
        flags=re.S
    )
    if block:
        for line in block.group(1).splitlines():
            line = line.strip()
            if not line:
                continue
            parts = re.split(r"\s+", line)
            if len(parts) >= 2:
                try:
                    cell_distances.append(float(parts[1]))
                except Exception:
                    pass

    # Transformation matrix (3x3)
    mat_m = re.search(
        r"Transformation matrix\s+([0-9\.\-\s]+)\n\s*([0-9\.\-\s]+)\n\s*([0-9\.\-\s]+)",
        text
    )
    if not mat_m:
        raise ValueError("Kon 'Transformation matrix' niet vinden in .hdr.")
    rows = []
    for i in range(1, 4):
        row = [float(x) for x in mat_m.group(i).split()]
        if len(row) != 3:
            raise ValueError("Transformation matrix rij is niet 3 waarden breed.")
        rows.append(row)
    T = np.array(rows, dtype=float)

    pressure_column = _extract_sen_pressure_column(text)
    if n_meas is None or n_cells is None or dt_sec is None:
        raise ValueError("Kon n_meas / n_cells / Profile interval niet volledig vinden in .hdr.")

    return {
        "n_meas": n_meas,
        "n_cells": n_cells,
        "dt_sec": dt_sec,
        "t0": t0,
        "cell_distances_m": cell_distances,
        "T": T,
        "sen_pressure_column": pressure_column,
    }


def build_time_index(t0: pd.Timestamp, n_meas: int, dt_sec: float) -> pd.DatetimeIndex:
    freq = pd.to_timedelta(dt_sec, unit="s")
    return pd.date_range(start=t0, periods=n_meas, freq=freq)


def load_v_file(path: str, n_meas: int, n_cells: int) -> np.ndarray:
    arr = np.loadtxt(path, dtype=float)
    if arr.ndim == 1:
        arr = arr.reshape(1, -1)
    if arr.shape[1] != n_cells:
        raise ValueError(f"{os.path.basename(path)}: verwacht {n_cells} kolommen, kreeg {arr.shape[1]}.")
    if arr.shape[0] != n_meas:
        min_n = min(arr.shape[0], n_meas)
        arr = arr[:min_n, :]
    return arr


def load_csv_processed(csv_path: str) -> pd.DataFrame:
    df = pd.read_csv(csv_path, sep=";", engine="python")
    if df.columns.size > 0 and (df.columns[-1].strip() == "" or "Unnamed" in df.columns[-1]):
        df = df.iloc[:, :-1]
    df.columns = [c.strip() for c in df.columns]

    if "DateTime" not in df.columns:
        raise ValueError("CSV mist kolom 'DateTime'.")

    df["DateTime"] = pd.to_datetime(
        df["DateTime"].astype(str).str.strip(),
        format="%d/%m/%Y %H:%M:%S",
        dayfirst=True,
        errors="coerce"
    )

    for c in df.columns:
        if c == "DateTime":
            continue
        df[c] = pd.to_numeric(
            df[c].astype(str).str.replace(",", ".", regex=False).str.strip(),
            errors="coerce"
        )

    return df


def load_sen_pressure(sen_path: str, pressure_column: int) -> pd.Series:
    if pressure_column < 8:
        raise ValueError("Pressure column index moet >= 8 zijn.")
    times = []
    values = []
    with open(sen_path, "r", errors="ignore") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = re.split(r"\s+", line)
            if len(parts) < 7:
                continue
            try:
                month, day, year, hour, minute, second = map(int, parts[:6])
            except ValueError:
                continue
            try:
                dt = pd.Timestamp(year, month, day, hour, minute, second)
            except ValueError:
                continue
            data_values = parts[6:]
            idx = pressure_column - 7
            val = np.nan
            if 0 <= idx < len(data_values):
                raw = data_values[idx].strip().replace(",", ".")
                if raw:
                    try:
                        val = float(raw)
                    except ValueError:
                        val = np.nan
            times.append(dt)
            values.append(val)
    if not times:
        return pd.Series([], dtype=float)
    series = pd.Series(values, index=pd.DatetimeIndex(times))
    return series


def speed_dir_to_uv(speed: np.ndarray, direction_deg: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    th = np.deg2rad(direction_deg)
    u = speed * np.sin(th)
    v = speed * np.cos(th)
    return u, v


def uv_to_speed_dir(u: np.ndarray, v: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    speed = np.sqrt(u*u + v*v)
    direction = (np.rad2deg(np.arctan2(u, v)) + 360.0) % 360.0
    return speed, direction


def compute_xyz_from_beams(T: np.ndarray, b1: np.ndarray, b2: np.ndarray, b3: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    B = np.stack([b1, b2, b3], axis=-1)
    XYZ = np.einsum("ij,tcj->tci", T, B)
    x = XYZ[..., 0]
    y = XYZ[..., 1]
    z = XYZ[..., 2]
    return x, y, z


# ----------------------------
# Plot helpers
# ----------------------------
def _new_or_clear_figure(fig_key: str, registry: dict, title: str):
    if fig_key in registry and plt.fignum_exists(registry[fig_key].number):
        fig = registry[fig_key]
        fig.clf()
    else:
        fig = plt.figure()
        registry[fig_key] = fig
    fig.suptitle(title)
    return fig


def _format_time_axis(ax):
    locator = mdates.AutoDateLocator()
    formatter = mdates.ConciseDateFormatter(locator)
    ax.xaxis.set_major_locator(locator)
    ax.xaxis.set_major_formatter(formatter)
    ax.tick_params(axis="x", rotation=90)


def _apply_tight_layout(fig):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        fig.tight_layout()


def _plot_pressure_panel(ax, time_index, pressure_arr):
    """
    Tekent de druk (pressure) onder een tijdreeks.
    """
    ax.set_ylabel("Pressure (dbar)")
    ax.grid(True, alpha=0.2)
    if pressure_arr is None or len(pressure_arr) == 0 or np.all(np.isnan(pressure_arr)):
        ax.text(0.5, 0.5, "Geen drukgegevens beschikbaar.", ha="center", va="center", color="gray",
                transform=ax.transAxes)
        ax.set_yticks([])
        _format_time_axis(ax)
        return

    ax.plot(time_index, pressure_arr, color="tab:cyan", linewidth=1.25)
    _format_time_axis(ax)
    ax.set_xlabel("Tijd")


def plot_beam(time_index, beam_arr, selected_cells, cell_distances, beam_name, fig_registry, pressure_arr=None):
    fig = _new_or_clear_figure(f"beam_{beam_name}", fig_registry, f"{beam_name} - snelheden per cel + gemiddelde")
    gs = fig.add_gridspec(2, 1, height_ratios=[3, 1], hspace=0.25)
    ax = fig.add_subplot(gs[0, 0])
    ax_pressure = fig.add_subplot(gs[1, 0], sharex=ax)
    ax.tick_params(labelbottom=False)

    if not selected_cells:
        ax.text(0.5, 0.5, "Geen cellen geselecteerd.", ha="center", va="center", transform=ax.transAxes)
        _apply_tight_layout(fig)
        fig.show()
        return

    for ci in selected_cells:
        label = f"Cel {ci+1}"
        if cell_distances and ci < len(cell_distances):
            label += f" ({cell_distances[ci]:.2f} m)"
        ax.plot(time_index, beam_arr[:, ci], label=label, linewidth=0.9)

    avg = np.nanmean(beam_arr[:, selected_cells], axis=1)
    ax.plot(time_index, avg, label="Gemiddelde (geselecteerde cellen)", linewidth=2.0)

    ax.set_ylabel("Snelheid langs beam (m/s)")
    ax.grid(True, alpha=0.3)
    ax.legend(loc="upper right", fontsize=8)
    _plot_pressure_panel(ax_pressure, time_index, pressure_arr)
    _apply_tight_layout(fig)
    fig.show()


def plot_resultant_from_csv(df_csv, selected_cells, fig_registry, pressure_arr=None):
    """
    Berekent resultante speed+dir over geselecteerde cellen vanuit Speed#i / Dir#i in CSV.
    """
    fig = _new_or_clear_figure("resultant_csv", fig_registry, "Resultante snelheid + richting (uit CSV, over geselecteerde cellen)")
    gs = fig.add_gridspec(2, 1, height_ratios=[3, 1], hspace=0.25)
    ax1 = fig.add_subplot(gs[0, 0])
    ax_pressure = fig.add_subplot(gs[1, 0], sharex=ax1)
    ax1.tick_params(labelbottom=False)

    if df_csv is None or df_csv.empty:
        ax1.text(0.5, 0.5, "CSV niet ingelezen.", ha="center", va="center", transform=ax1.transAxes)
        _apply_tight_layout(fig)
        _plot_pressure_panel(ax_pressure, pd.DatetimeIndex([]), None)
        fig.show()
        return

    if not selected_cells:
        ax1.text(0.5, 0.5, "Geen cellen geselecteerd.", ha="center", va="center", transform=ax1.transAxes)
        _apply_tight_layout(fig)
        _plot_pressure_panel(ax_pressure, pd.DatetimeIndex([]), None)
        fig.show()
        return

    u_list = []
    v_list = []
    for ci in selected_cells:
        k = ci + 1
        speed_col = [c for c in df_csv.columns if c.strip().startswith(f"Speed#{k}")]
        dir_col = [c for c in df_csv.columns if c.strip().startswith(f"Dir#{k}")]
        if not speed_col or not dir_col:
            continue
        sp = df_csv[speed_col[0]].to_numpy(dtype=float)
        dr = df_csv[dir_col[0]].to_numpy(dtype=float)
        u, v = speed_dir_to_uv(sp, dr)
        u_list.append(u)
        v_list.append(v)

    if not u_list:
        ax1.text(0.5, 0.5, "Geen Speed#/Dir# kolommen gevonden voor selectie.", ha="center", va="center", transform=ax1.transAxes)
        _apply_tight_layout(fig)
        _plot_pressure_panel(ax_pressure, pd.DatetimeIndex([]), None)
        fig.show()
        return

    U = np.nanmean(np.vstack(u_list), axis=0)
    V = np.nanmean(np.vstack(v_list), axis=0)
    spd, direc = uv_to_speed_dir(U, V)

    t = df_csv["DateTime"]
    ax1.plot(t, spd, linewidth=1.5)
    ax1.set_ylabel("Resultante snelheid (m/s)")
    ax1.grid(True, alpha=0.3)

    ax2 = ax1.twinx()
    ax2.plot(t, direc, linewidth=1.0)
    ax2.set_ylabel("Richting (deg)")

    _plot_pressure_panel(ax_pressure, t, pressure_arr)
    _apply_tight_layout(fig)
    fig.show()
def plot_component(time_index, comp_arr, selected_cells, cell_distances, title, ylabel, fig_key, fig_registry, pressure_arr=None):
    fig = _new_or_clear_figure(fig_key, fig_registry, title)
    gs = fig.add_gridspec(2, 1, height_ratios=[3, 1], hspace=0.25)
    ax = fig.add_subplot(gs[0, 0])
    ax_pressure = fig.add_subplot(gs[1, 0], sharex=ax)
    ax.tick_params(labelbottom=False)

    if not selected_cells:
        ax.text(0.5, 0.5, "Geen cellen geselecteerd.", ha="center", va="center", transform=ax.transAxes)
        _apply_tight_layout(fig)
        fig.show()
        return

    for ci in selected_cells:
        label = f"Cel {ci+1}"
        if cell_distances and ci < len(cell_distances):
            label += f" ({cell_distances[ci]:.2f} m)"
        ax.plot(time_index, comp_arr[:, ci], label=label, linewidth=0.9)

    avg = np.nanmean(comp_arr[:, selected_cells], axis=1)
    ax.plot(time_index, avg, label="Gemiddelde (geselecteerde cellen)", linewidth=2.0)

    ax.set_ylabel(ylabel)
    ax.grid(True, alpha=0.3)
    ax.legend(loc="upper right", fontsize=8)
    _plot_pressure_panel(ax_pressure, time_index, pressure_arr)
    _apply_tight_layout(fig)
    fig.show()
def apply_running_mean(arr: np.ndarray, window: int) -> np.ndarray:
    """
    Maakt een lopend gemiddelde over tijd (axis=0) voor alle cellen.
    """
    if window <= 1:
        return arr
    arr_np = np.asarray(arr)
    if arr_np.ndim == 1:
        sr = pd.Series(arr_np)
        return sr.rolling(window=window, min_periods=1, center=True).mean().to_numpy()
    df = pd.DataFrame(arr_np)
    return df.rolling(window=window, min_periods=1, center=True).mean().to_numpy()


def plot_current_rose(u: np.ndarray,
                      v: np.ndarray,
                      title: str,
                      fig_key: str,
                      fig_registry: dict,
                      n_dir_bins: int = 36):
    """
    Stroomroos (current rose) op basis van u,v componenten.
    Richting = waar de stroming NAARTOE gaat (oceanografisch).
    """

    speed = np.sqrt(u**2 + v**2)
    direction = (np.degrees(np.arctan2(u, v)) + 360.0) % 360.0

    bins = np.linspace(0, 360, n_dir_bins + 1)
    bin_centers = 0.5 * (bins[:-1] + bins[1:])
    bin_width = 2 * np.pi / n_dir_bins

    mean_speed = np.zeros(n_dir_bins)
    for i in range(n_dir_bins):
        mask = (direction >= bins[i]) & (direction < bins[i + 1])
        if np.any(mask):
            mean_speed[i] = np.nanmean(speed[mask])

    theta = np.deg2rad(bin_centers)

    if fig_key in fig_registry and plt.fignum_exists(fig_registry[fig_key].number):
        fig = fig_registry[fig_key]
        fig.clf()
    else:
        fig = plt.figure()
        fig_registry[fig_key] = fig

    ax = fig.add_subplot(111, polar=True)
    ax.bar(theta, mean_speed, width=bin_width, bottom=0.0)

    ax.set_theta_zero_location("N")
    ax.set_theta_direction(-1)
    ax.set_title(title)

    _apply_tight_layout(fig)
    fig.show()


# ----------------------------
# Export helper (EU formatting)
# ----------------------------
def export_tsv_eu(df: pd.DataFrame, out_path: str, decimals: int = 6):
    """
    Exporteer tab-separated met:
      - decimaal = ','
      - duizendtallen = '.'
    """
    df_out = df.copy()

    for col in df_out.columns:
        if pd.api.types.is_numeric_dtype(df_out[col]):
            df_out[col] = df_out[col].map(
                lambda x: (
                    f"{x:,.{decimals}f}"
                    .replace(",", "TMP")
                    .replace(".", ",")
                    .replace("TMP", ".")
                ) if pd.notnull(x) else ""
            )

    df_out.to_csv(out_path, sep="\t", index=False)


def filter_time_window(time_index: pd.DatetimeIndex, start_dt: pd.Timestamp | None, end_dt: pd.Timestamp | None):
    """
    Returns boolean mask for time_index within [start_dt, end_dt] (inclusive).
    """
    mask = np.ones(len(time_index), dtype=bool)
    if start_dt is not None:
        mask &= (time_index >= start_dt)
    if end_dt is not None:
        mask &= (time_index <= end_dt)
    return mask


# ----------------------------
# GUI App
# ----------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Aquadopp - beams (v1/v2/v3) + CSV controle - GUI")
        self.geometry("1040x700")

        self.meta = None
        self.base_hdr = None
        self.base_prefix = None

        self.v1 = None
        self.v2 = None
        self.v3 = None
        self.df_csv = None
        self.time_index = None
        self.pressure_series = None

        self.fig_registry = {}

        # UI state
        self.cell_vars = []
        self.export_var = tk.BooleanVar(value=False)
        self.running_mean_var = tk.StringVar(value="1")
        self.rose_mode_var = tk.StringVar(value="XY")

        self.start_var = tk.StringVar(value="")
        self.end_var = tk.StringVar(value="")

        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self, padding=10)
        top.pack(side=tk.TOP, fill=tk.X)

        btn_pick = ttk.Button(top, text="Kies .hdr bestand...", command=self.pick_hdr)
        btn_pick.pack(side=tk.LEFT)

        self.lbl_file = ttk.Label(top, text="(geen bestand gekozen)")
        self.lbl_file.pack(side=tk.LEFT, padx=10)

        mid = ttk.Frame(self, padding=10)
        mid.pack(side=tk.TOP, fill=tk.BOTH, expand=True)

        # left: cell selection (scrollable)
        left = ttk.LabelFrame(mid, text="Cellen (aanvinken)", padding=8)
        left.pack(side=tk.LEFT, fill=tk.Y)

        self.canvas = tk.Canvas(left, width=320, height=500)
        self.scrollbar = ttk.Scrollbar(left, orient="vertical", command=self.canvas.yview)
        self.scroll_frame = ttk.Frame(self.canvas)

        self.scroll_frame.bind(
            "<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )
        self.canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # right: actions/info
        right = ttk.Frame(mid, padding=8)
        right.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        # Time window frame
        tw = ttk.LabelFrame(right, text="Tijdvenster (zoom / export)", padding=8)
        tw.pack(fill=tk.X)

        ttk.Label(tw, text="Begin (dd/mm/yyyy hh:mm:ss) leeg = start").grid(row=0, column=0, sticky="w")
        e1 = ttk.Entry(tw, textvariable=self.start_var, width=28)
        e1.grid(row=0, column=1, padx=6, pady=2, sticky="w")

        ttk.Label(tw, text="Einde (dd/mm/yyyy hh:mm:ss) leeg = eind").grid(row=1, column=0, sticky="w")
        e2 = ttk.Entry(tw, textvariable=self.end_var, width=28)
        e2.grid(row=1, column=1, padx=6, pady=2, sticky="w")

        btn_set_full = ttk.Button(tw, text="Vul volledig bereik in", command=self.fill_full_range)
        btn_set_full.grid(row=0, column=2, rowspan=2, padx=10, sticky="ns")

        tw.columnconfigure(3, weight=1)

        self.txt_info = tk.Text(right, height=12, wrap="word")
        self.txt_info.pack(fill=tk.X, pady=(10, 0))
        self.txt_info.insert("end", "Kies een .hdr bestand om te starten.\n")
        self.txt_info.configure(state="disabled")

        action = ttk.Frame(right)
        action.pack(fill=tk.X, pady=10)

        chk_export = ttk.Checkbutton(action, text="Export bij herbereken", variable=self.export_var)
        chk_export.pack(side=tk.LEFT)

        btn_run = ttk.Button(action, text="Herbereken en teken", command=self.recompute_and_plot)
        btn_run.pack(side=tk.LEFT, padx=10)

        btn_save_npz = ttk.Button(action, text="Schrijf NPZ weg (beams+xyz)", command=self.save_npz)
        btn_save_npz.pack(side=tk.LEFT)

        btn_export_word = ttk.Button(action, text="Exporteer figuren -> Word", command=self.export_figures_to_word)
        btn_export_word.pack(side=tk.LEFT, padx=(6, 0))

        opts = ttk.Frame(right)
        opts.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(opts, text="Stroomroos uit:").pack(side=tk.LEFT)
        cmb = ttk.Combobox(
            opts,
            textvariable=self.rose_mode_var,
            state="readonly",
            values=["XY", "Alleen X"],
            width=10
        )
        cmb.pack(side=tk.LEFT, padx=6)

        ttk.Label(opts, text="Lopend gemiddelde (N):").pack(side=tk.LEFT, padx=(16, 0))
        ent_rm = ttk.Entry(opts, textvariable=self.running_mean_var, width=6)
        ent_rm.pack(side=tk.LEFT, padx=(4, 0))

        tips = ttk.Label(right, text=(
            "Opmerking:\n"
            "* Tijdas wordt opgebouwd uit starttijd + Profile interval.\n"
            "* XYZ via Transformation matrix uit .hdr: [x;y;z]=T*[b1; b2; b3].\n"
            "* Resultante speed/dir uit CSV Speed#/Dir# over geselecteerde cellen.\n"
            "* Tijdvenster geldt voor plots en export.\n"
            "* Kies stroomroos-bron en eventueel lopend gemiddelde in dezelfde actie-regel."
        ))
        tips.pack(fill=tk.X, pady=10)

    def log(self, msg: str):
        self.txt_info.configure(state="normal")
        self.txt_info.insert("end", msg.rstrip() + "\n")
        self.txt_info.see("end")
        self.txt_info.configure(state="disabled")

    def parse_time_entry(self, s: str) -> pd.Timestamp | None:
        s = (s or "").strip()
        if not s:
            return None
        try:
            return pd.to_datetime(s, format="%d/%m/%Y %H:%M:%S", dayfirst=True)
        except Exception:
            raise ValueError(f"Ongeldig tijdformaat: '{s}'. Verwacht dd/mm/yyyy hh:mm:ss")

    def _align_pressure_series(self, pressure_series: pd.Series, time_index: pd.DatetimeIndex, dt_sec: float) -> pd.Series | None:
        if pressure_series is None or pressure_series.empty:
            return None
        df_base = pd.DataFrame({"DateTime": time_index})
        df_pressure = pressure_series.dropna().sort_index().reset_index()
        df_pressure.columns = ["DateTime", "pressure"]
        tolerance = pd.Timedelta(seconds=max(dt_sec, 1.0))
        merged = pd.merge_asof(df_base, df_pressure, on="DateTime", direction="nearest", tolerance=tolerance)
        merged = merged.sort_values("DateTime")
        return pd.Series(merged["pressure"].to_numpy(), index=time_index)


    def fill_full_range(self):
        if self.time_index is None or len(self.time_index) == 0:
            return
        self.start_var.set(self.time_index[0].strftime("%d/%m/%Y %H:%M:%S"))
        self.end_var.set(self.time_index[-1].strftime("%d/%m/%Y %H:%M:%S"))

    def _get_running_mean_window(self) -> int:
        value = (self.running_mean_var.get() or "").strip()
        if not value:
            return 1
        try:
            window = int(value)
        except ValueError:
            raise ValueError("Lopend gemiddelde moet een positief geheel getal zijn.")
        if window < 1:
            raise ValueError("Lopend gemiddelde moet minimaal 1 zijn.")
        return window

    def pick_hdr(self):
        path = filedialog.askopenfilename(
            title="Kies .hdr bestand",
            filetypes=[("HDR files", "*.hdr"), ("All files", "*.*")]
        )
        if not path:
            return

        try:
            meta = parse_hdr(path)
        except Exception as e:
            messagebox.showerror("Fout", f"Kan .hdr niet parsen:\n{e}")
            return

        base_prefix = os.path.splitext(path)[0]
        v1_path = base_prefix + ".v1"
        v2_path = base_prefix + ".v2"
        v3_path = base_prefix + ".v3"
        csv_path = base_prefix + ".csv"

        missing = [p for p in [v1_path, v2_path, v3_path, csv_path] if not os.path.exists(p)]
        if missing:
            messagebox.showwarning(
                "Bestanden ontbreken",
                "Volgende bestanden ontbreken:\n\n" + "\n".join(missing) +
                "\n\nIk ga verder met wat beschikbaar is."
            )

        # load v-files
        try:
            self.v1 = load_v_file(v1_path, meta["n_meas"], meta["n_cells"]) if os.path.exists(v1_path) else None
            self.v2 = load_v_file(v2_path, meta["n_meas"], meta["n_cells"]) if os.path.exists(v2_path) else None
            self.v3 = load_v_file(v3_path, meta["n_meas"], meta["n_cells"]) if os.path.exists(v3_path) else None
        except Exception as e:
            messagebox.showerror("Fout", f"Kan v-bestanden niet inlezen:\n{e}")
            return

        # time index length = min beschikbare arrays
        n_meas_eff = meta["n_meas"]
        for arr in [self.v1, self.v2, self.v3]:
            if arr is not None:
                n_meas_eff = min(n_meas_eff, arr.shape[0])
        self.time_index = build_time_index(meta["t0"], n_meas_eff, meta["dt_sec"])

        # trim v arrays
        if self.v1 is not None:
            self.v1 = self.v1[:n_meas_eff, :]
        if self.v2 is not None:
            self.v2 = self.v2[:n_meas_eff, :]
        if self.v3 is not None:
            self.v3 = self.v3[:n_meas_eff, :]


        self.pressure_series = None
        sen_path = base_prefix + '.sen'
        if os.path.exists(sen_path):
            try:
                raw_pressure = load_sen_pressure(sen_path, meta.get("sen_pressure_column", 14))
                aligned = self._align_pressure_series(raw_pressure, self.time_index, meta["dt_sec"])
                if aligned is not None:
                    self.pressure_series = aligned
                    self.log(f"Druk uit .sen geladen: {len(aligned)} punten")
                else:
                    self.log("Drukgegevens konden niet op tijdsindex worden afgestemd.")
            except Exception as exc:
                self.log(f"Kon drukwaarden uit .sen niet inlezen: {exc}")
        else:
            self.log(".sen bestand niet gevonden -> drukplots worden overgeslagen.")
        # load csv
        self.df_csv = None
        if os.path.exists(csv_path):
            try:
                self.df_csv = load_csv_processed(csv_path)
            except Exception as e:
                messagebox.showwarning("CSV probleem", f"CSV kon niet goed ingelezen worden:\n{e}")

        self.meta = meta
        self.base_hdr = path
        self.base_prefix = base_prefix

        self.lbl_file.configure(text=os.path.basename(path))
        self.log(f"Gekozen: {path}")
        self.log(f"n_meas={meta['n_meas']} | n_cells={meta['n_cells']} | dt={meta['dt_sec']} s | start={meta['t0']}")
        if meta["cell_distances_m"]:
            self.log("Cell center distances (m): " + ", ".join(f"{x:.2f}" for x in meta["cell_distances_m"]))
        self.log("Transformation matrix T:\n" + "\n".join("  " + "  ".join(f"{v:8.4f}" for v in r) for r in meta["T"]))

        # rebuild cell checkboxes
        for w in self.scroll_frame.winfo_children():
            w.destroy()
        self.cell_vars = []
        n_cells = meta["n_cells"]
        dists = meta["cell_distances_m"]

        for i in range(n_cells):
            v = tk.BooleanVar(value=True)
            self.cell_vars.append(v)
            label = f"Cel {i+1}"
            if dists and i < len(dists):
                label += f"  ({dists[i]:.2f} m)"
            cb = ttk.Checkbutton(self.scroll_frame, text=label, variable=v)
            cb.pack(anchor="w")

        # fill time range defaults
        self.fill_full_range()

        self.log("Cellen aangevinkt (standaard: allemaal). Tijdvenster ingevuld met volledig bereik.")

    def selected_cells(self):
        return [i for i, v in enumerate(self.cell_vars) if v.get()]

    def recompute_and_plot(self):
        if self.meta is None or self.time_index is None:
            messagebox.showinfo("Info", "Kies eerst een .hdr bestand.")
            return
        if self.v1 is None or self.v2 is None or self.v3 is None:
            messagebox.showerror("Fout", "Minstens Ã©Ã©n van v1/v2/v3 ontbreekt. XYZ en beamplots vereisen alle 3.")
            return

        try:
            rm_window = self._get_running_mean_window()
        except ValueError as exc:
            messagebox.showerror("Lopend gemiddelde fout", str(exc))
            return

        # parse window
        try:
            start_dt = self.parse_time_entry(self.start_var.get())
            end_dt = self.parse_time_entry(self.end_var.get())
        except Exception as e:
            messagebox.showerror("Tijdvenster fout", str(e))
            return

        if start_dt and end_dt and end_dt < start_dt:
            messagebox.showerror("Tijdvenster fout", "Eindmoment ligt vÃ³Ã³r beginmoment.")
            return

        mask = filter_time_window(self.time_index, start_dt, end_dt)
        if not mask.any():
            messagebox.showwarning("Tijdvenster", "Geen data binnen dit tijdvenster.")
            return

        ti = self.time_index[mask]
        pressure_plot = None
        if self.pressure_series is not None:
            pressure_window = np.asarray(self.pressure_series.iloc[mask])
            if pressure_window.size > 0:
                pressure_plot = apply_running_mean(pressure_window, rm_window)
        v1w = self.v1[mask, :]
        v2w = self.v2[mask, :]
        v3w = self.v3[mask, :]
        v1w_plot = apply_running_mean(v1w, rm_window)
        v2w_plot = apply_running_mean(v2w, rm_window)
        v3w_plot = apply_running_mean(v3w, rm_window)

        sel = self.selected_cells()
        dists = self.meta["cell_distances_m"]

        # Beam figures (windowed)
        plot_beam(ti, v1w_plot, sel, dists, "Beam 1 (v1)", self.fig_registry, pressure_arr=pressure_plot)
        plot_beam(ti, v2w_plot, sel, dists, "Beam 2 (v2)", self.fig_registry, pressure_arr=pressure_plot)
        plot_beam(ti, v3w_plot, sel, dists, "Beam 3 (v3)", self.fig_registry, pressure_arr=pressure_plot)

        # Resultant from CSV within window (filter by DateTime)
        if self.df_csv is not None:
            dfw = self.df_csv.copy()
            if start_dt is not None:
                dfw = dfw[dfw["DateTime"] >= start_dt]
            if end_dt is not None:
                dfw = dfw[dfw["DateTime"] <= end_dt]
            plot_resultant_from_csv(dfw, sel, self.fig_registry, pressure_arr=pressure_plot)
        else:
            self.log("Geen CSV geladen -> resultante speed/dir figuur wordt overgeslagen.")

        # XYZ via transformation matrix (windowed)
        T = self.meta["T"]
        x, y, z = compute_xyz_from_beams(T, v1w, v2w, v3w)
        x_plot = apply_running_mean(x, rm_window)
        y_plot = apply_running_mean(y, rm_window)
        z_plot = apply_running_mean(z, rm_window)

        plot_component(ti, x_plot, sel, dists, "X-projectie (via T*[b1 b2 b3])", "X (m/s)", "comp_x", self.fig_registry, pressure_arr=pressure_plot)
        plot_component(ti, y_plot, sel, dists, "Y-projectie (via T*[b1 b2 b3])", "Y (m/s)", "comp_y", self.fig_registry, pressure_arr=pressure_plot)
        plot_component(ti, z_plot, sel, dists, "Z-projectie (via T*[b1 b2 b3])", "Z (m/s)", "comp_z", self.fig_registry, pressure_arr=pressure_plot)

        if sel:
            u_sel = np.nanmean(x_plot[:, sel], axis=1)
            if self.rose_mode_var.get() == "Alleen X":
                v_sel = np.zeros_like(u_sel)
            else:
                v_sel = np.nanmean(y_plot[:, sel], axis=1)

            plot_current_rose(
                u=u_sel,
                v=v_sel,
                title="Stroomroos - gemiddelde stroming (gekozen cellen, tijdvenster)",
                fig_key="current_rose",
                fig_registry=self.fig_registry,
            )

        # Optional export (windowed)
        if self.export_var.get():
            try:
                self._export_current_windowed(sel, ti, v1w_plot, v2w_plot, v3w_plot, x_plot, y_plot, z_plot, start_dt, end_dt)
            except Exception as e:
                messagebox.showwarning("Export fout", f"Export faalde:\n{e}")

    def _export_current_windowed(self, sel, ti, v1w, v2w, v3w, x, y, z, start_dt, end_dt):
        out_default = self.base_prefix + "_export_window.tsv"
        out_path = filedialog.asksaveasfilename(
            title="Kies exportbestand",
            defaultextension=".tsv",
            initialfile=os.path.basename(out_default),
            filetypes=[("TSV (tab separated)", "*.tsv"), ("All files", "*.*")]
        )
        if not out_path:
            return

        b1_avg = np.nanmean(v1w[:, sel], axis=1) if sel else np.full(len(ti), np.nan)
        b2_avg = np.nanmean(v2w[:, sel], axis=1) if sel else np.full(len(ti), np.nan)
        b3_avg = np.nanmean(v3w[:, sel], axis=1) if sel else np.full(len(ti), np.nan)

        x_avg = np.nanmean(x[:, sel], axis=1) if sel else np.full(len(ti), np.nan)
        y_avg = np.nanmean(y[:, sel], axis=1) if sel else np.full(len(ti), np.nan)
        z_avg = np.nanmean(z[:, sel], axis=1) if sel else np.full(len(ti), np.nan)

        df_out = pd.DataFrame({
            "DateTime": ti,
            "beam1_avg": b1_avg,
            "beam2_avg": b2_avg,
            "beam3_avg": b3_avg,
            "x_avg": x_avg,
            "y_avg": y_avg,
            "z_avg": z_avg,
        })

        # CSV toevoegen indien beschikbaar (filtered window)
        if self.df_csv is not None and "DateTime" in self.df_csv.columns:
            df_csv = self.df_csv.copy().sort_values("DateTime")
            if start_dt is not None:
                df_csv = df_csv[df_csv["DateTime"] >= start_dt]
            if end_dt is not None:
                df_csv = df_csv[df_csv["DateTime"] <= end_dt]

            # resultante uit csv over selectie
            u_list, v_list = [], []
            for ci in sel:
                k = ci + 1
                speed_col = [c for c in df_csv.columns if c.strip().startswith(f"Speed#{k}")]
                dir_col = [c for c in df_csv.columns if c.strip().startswith(f"Dir#{k}")]
                if not speed_col or not dir_col:
                    continue
                sp = df_csv[speed_col[0]].to_numpy(dtype=float)
                dr = df_csv[dir_col[0]].to_numpy(dtype=float)
                u, v = speed_dir_to_uv(sp, dr)
                u_list.append(u)
                v_list.append(v)

            if u_list:
                U = np.nanmean(np.vstack(u_list), axis=0)
                V = np.nanmean(np.vstack(v_list), axis=0)
                spd, direc = uv_to_speed_dir(U, V)
                df_csv = df_csv.copy()
                df_csv["resultant_speed_sel"] = spd
                df_csv["resultant_dir_sel"] = direc

            keep_cols = [c for c in ["Battery", "Heading", "Pitch", "Roll", "Pressure", "Temperature",
                                     "resultant_speed_sel", "resultant_dir_sel"] if c in df_csv.columns]
            df_csv_small = df_csv[["DateTime"] + keep_cols].copy()

            df_out = df_out.sort_values("DateTime")
            df_csv_small = df_csv_small.sort_values("DateTime")
            df_out = pd.merge_asof(df_out, df_csv_small, on="DateTime", direction="nearest")

        export_tsv_eu(df_out, out_path)
        self.log(f"Export geschreven: {out_path}")

    def export_figures_to_word(self):
        if Document is None or Inches is None:
            messagebox.showerror(
                "Fout",
                "De module 'python-docx' ontbreekt. Installeer deze via 'pip install python-docx' om grafieken naar Word te exporteren."
            )
            return

        figs = [
            (name, fig)
            for name, fig in self.fig_registry.items()
            if plt.fignum_exists(fig.number)
        ]
        if not figs:
            messagebox.showinfo("Info", "Er zijn geen actieve figuren om te exporteren.")
            return

        temp_dir = tempfile.mkdtemp(prefix="aqp_figs_")
        saved = []
        try:
            for name, fig in figs:
                sanitized = name.replace(" ", "_")
                path = os.path.join(temp_dir, f"{sanitized}.png")
                fig.savefig(path, bbox_inches="tight")
                saved.append((name, path))

            doc = Document()
            doc.add_heading("Aquadopp - Openstaande grafieken", level=1)
            start_str = (self.start_var.get() or "").strip()
            end_str = (self.end_var.get() or "").strip()
            if not start_str:
                start_str = "begin (volledig bereik)"
            if not end_str:
                end_str = "einde (volledig bereik)"
            window_desc = (
                f"Tijdvenster: {start_str} -> {end_str}\n"
                f"Lopend gemiddelde: {self.running_mean_var.get() or '1'}\n"
                f"Stroomroos-bron: {self.rose_mode_var.get()}\n"
                f"Geselecteerde cellen: {len(self.selected_cells())}"
            )
            doc.add_paragraph("Instellingen", style="Heading 2")
            for line in window_desc.splitlines():
                doc.add_paragraph(line, style="List Bullet")
            for idx, (name, path) in enumerate(saved):
                if idx:
                    doc.add_page_break()
                doc.add_paragraph(name)
                doc.add_picture(path, width=Inches(6))

            initialdir = os.path.dirname(self.base_prefix) if self.base_prefix else os.getcwd()
            initialfile = (
                f"{os.path.basename(self.base_prefix) if self.base_prefix else 'aquadopp'}_figuren.docx"
            )
            out_path = filedialog.asksaveasfilename(
                title="Kies Word-bestand",
                initialdir=initialdir,
                initialfile=initialfile,
                defaultextension=".docx",
                filetypes=[("Word document", "*.docx"), ("All files", "*.*")]
            )
            if not out_path:
                return
            doc.save(out_path)
            self.log(f"Figuren in Word geschreven: {out_path}")
        except Exception as exc:
            messagebox.showerror("Fout", f"Kon Word-export niet uitvoeren:\n{exc}")
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def save_npz(self):
        if self.meta is None or self.time_index is None:
            messagebox.showinfo("Info", "Kies eerst een .hdr bestand.")
            return
        if self.v1 is None or self.v2 is None or self.v3 is None:
            messagebox.showerror("Fout", "v1/v2/v3 vereist.")
            return

        out_default = self.base_prefix + "_beams_xyz.npz"
        out_path = filedialog.asksaveasfilename(
            title="Kies NPZ opslagbestand",
            defaultextension=".npz",
            initialfile=os.path.basename(out_default),
            filetypes=[("NPZ", "*.npz"), ("All files", "*.*")]
        )
        if not out_path:
            return

        T = self.meta["T"]
        x, y, z = compute_xyz_from_beams(T, self.v1, self.v2, self.v3)

        np.savez_compressed(
            out_path,
            time=self.time_index.astype("datetime64[ns]").values,
            cell_distances_m=np.array(self.meta["cell_distances_m"], dtype=float) if self.meta["cell_distances_m"] else np.array([]),
            T=T,
            v1=self.v1,
            v2=self.v2,
            v3=self.v3,
            x=x, y=y, z=z
        )
        self.log(f"NPZ geschreven: {out_path}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
